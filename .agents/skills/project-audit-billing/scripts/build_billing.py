#!/usr/bin/env python3
"""Build a period invoice summary (markdown) and billing workbook from billing/config.json + billing/ledger.csv.

Usage: python build_billing.py <project_root> <period_end YYYY-MM-DD> <xlsx_out>
Ledger columns: date,period,type,category,description,changelog_ref,hours,amount
type: work | retainer | platform_cost   (platform_cost is included in the retainer, never added on top)
"""
import csv, json, sys, datetime as dt
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference

root, end, xlsx_out = Path(sys.argv[1]), sys.argv[2], sys.argv[3]
cfg = json.loads((root / "billing/config.json").read_text())
rows = list(csv.DictReader(open(root / "billing/ledger.csv", newline="")))
day = int(cfg.get("billing_day", 18))
e = dt.date.fromisoformat(end)
pm = (e.replace(day=1) - dt.timedelta(days=1))
start = pm.replace(day=min(day, 28)) + dt.timedelta(days=1)
period = f"{start.isoformat()}..{e.isoformat()}"
cur = [r for r in rows if r["period"] == period]
work = [r for r in cur if r["type"] == "work"]
retainer = float(cfg["retainer"]); plat = float(cfg["platform_cost"])

# ---- markdown invoice backup
by_cat = defaultdict(list)
for r in work: by_cat[r["category"]].append(r)
md = [f"# Invoice backup — {cfg['client']}", "",
      f"**Billing period:** {start:%b %d, %Y} – {e:%b %d, %Y}  ", f"**Invoice date:** {e:%b %d, %Y}", "",
      "## Summary", "", "| Line | Amount |", "| --- | ---: |",
      f"| Monthly retainer | ${retainer:,.2f} |",
      f"| — includes: {cfg['platform_name']} platform cost | ${plat:,.2f} |",
      f"| — includes: design, development & support services | ${retainer - plat:,.2f} |",
      f"| **Total due** | **${retainer:,.2f}** |", ""]
if cfg.get("platform_cost_placeholder"):
    md += ["> Platform cost is a placeholder — confirm the actual amount.", ""]
md += [f"## Work completed ({len(work)} items)", ""]
for cat in sorted(by_cat):
    md += [f"### {cat}", ""]
    for r in sorted(by_cat[cat], key=lambda x: x["date"]):
        h = f" · {r['hours']}h" if r["hours"] else ""
        md.append(f"- {r['date']} — {r['description']} ({r['changelog_ref']}{h})")
    md.append("")
if cfg.get("open_items"):
    md += ["## Waiting on client", ""] + [f"- {i}" for i in cfg["open_items"]] + [""]
md += ["Full detail: see CHANGELOG.md and billing/ledger.csv."]
out = root / f"billing/invoices/{e.isoformat()}.md"
out.parent.mkdir(parents=True, exist_ok=True); out.write_text("\n".join(md) + "\n")

# ---- workbook
wb = Workbook(); hdr = Font(bold=True, color="FFFFFF", name="Arial"); fill = PatternFill("solid", start_color="5B2A55")
def head(ws, cols):
    ws.append(cols)
    for c in ws[1]: c.font, c.fill = hdr, fill
ws = wb.active; ws.title = "Ledger"
cols = ["date", "period", "type", "category", "description", "changelog_ref", "hours", "amount"]
head(ws, cols)
for r in rows:
    ws.append([r["date"], r["period"], r["type"], r["category"], r["description"], r["changelog_ref"],
               float(r["hours"]) if r["hours"] else None, float(r["amount"]) if r["amount"] else None])
n = ws.max_row
for w, col in zip([12, 24, 14, 14, 70, 14, 8, 10], "ABCDEFGH"): ws.column_dimensions[col].width = w
for c in ws["H"][1:]: c.number_format = "$#,##0.00;($#,##0.00);-"
if cfg.get("platform_cost_placeholder"):
    for row in ws.iter_rows(min_row=2):
        if row[2].value == "platform_cost": row[7].fill = PatternFill("solid", start_color="FFFF00")

periods = sorted({r["period"] for r in rows}); cats = sorted({r["category"] for r in rows if r["type"] == "work"})
ps = wb.create_sheet("Period Summary"); head(ps, ["period", "work items", "retainer", "platform cost (incl.)", "services portion"] + cats)
for i, p in enumerate(periods, start=2):
    ps.append([p, f'=COUNTIFS(Ledger!$B$2:$B${n},A{i},Ledger!$C$2:$C${n},"work")',
               f'=SUMIFS(Ledger!$H$2:$H${n},Ledger!$B$2:$B${n},A{i},Ledger!$C$2:$C${n},"retainer")',
               f'=SUMIFS(Ledger!$H$2:$H${n},Ledger!$B$2:$B${n},A{i},Ledger!$C$2:$C${n},"platform_cost")',
               f"=C{i}-D{i}"] +
              [f'=COUNTIFS(Ledger!$B$2:$B${n},$A{i},Ledger!$C$2:$C${n},"work",Ledger!$D$2:$D${n},"{c}")' for c in cats])
    for col in "CDE": ps[f"{col}{i}"].number_format = "$#,##0.00;($#,##0.00);-"
ps.column_dimensions["A"].width = 26
for col in "BCDE": ps.column_dimensions[col].width = 18

db = wb.create_sheet("Dashboard"); db["A1"] = f"{cfg['client']} — billing dashboard"; db["A1"].font = Font(bold=True, size=14, name="Arial")
last = len(periods) + 1
bc = BarChart(); bc.title = "Work items per period"; bc.y_axis.title = "items"
bc.add_data(Reference(ps, min_col=2, min_row=1, max_row=last), titles_from_data=True)
bc.set_categories(Reference(ps, min_col=1, min_row=2, max_row=last)); bc.width, bc.height = 18, 8
db.add_chart(bc, "A3")
db["L3"] = "category"; db["M3"] = "items (all time)"
for j, c in enumerate(cats, start=4):
    db[f"L{j}"] = c; db[f"M{j}"] = f'=COUNTIFS(Ledger!$C$2:$C${n},"work",Ledger!$D$2:$D${n},L{j})'
pc = PieChart(); pc.title = "Work by category"
pc.add_data(Reference(db, min_col=13, min_row=3, max_row=3 + len(cats)), titles_from_data=True)
pc.set_categories(Reference(db, min_col=12, min_row=4, max_row=3 + len(cats))); pc.width, pc.height = 14, 8
db.add_chart(pc, "A20")
wb.save(xlsx_out)
print(f"period {period}: {len(work)} work items -> {out}, {xlsx_out}")
